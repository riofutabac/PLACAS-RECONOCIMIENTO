"""Pruebas unitarias para el módulo lastre.excel."""

from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

from lastre.excel import ExcelError, escribir_listado


@pytest.fixture
def filas():
    return [
        {"registro": "V001", "placa": "PCW2497", "hora_paso": "16:17:47",
         "tiempo_video": "00:29.62", "video": "muestra60.mp4", "tipo_vehiculo": "automovil",
         "sentido": "sale", "confianza": "0.938", "lecturas": 4,
         "estado": "validado", "imagen": "placas/v01.jpg"},
        {"registro": "V002", "placa": "", "hora_paso": "16:19:04",
         "tiempo_video": "01:47.10", "video": "muestra60.mp4", "tipo_vehiculo": "camion",
         "sentido": "entra", "confianza": "0.000", "lecturas": 0,
         "estado": "sin placa identificable", "imagen": ""},
    ]


@pytest.fixture
def directorio_imagenes(tmp_path):
    """Crea una imagen real para probar la incrustación."""
    import cv2
    carpeta = tmp_path / "placas"
    carpeta.mkdir()
    cv2.imwrite(str(carpeta / "v01.jpg"), np.full((120, 300, 3), 200, dtype=np.uint8))
    return tmp_path


def test_genera_archivo_xlsx(tmp_path, filas, directorio_imagenes):
    """El listado se escribe como archivo de Excel."""
    destino = tmp_path / "placas.xlsx"

    resultado = escribir_listado(filas, destino, directorio_imagenes)

    assert resultado.is_file()
    assert resultado.suffix == ".xlsx"


def test_escribe_encabezado_y_una_fila_por_registro(tmp_path, filas, directorio_imagenes):
    """Cada vehículo ocupa exactamente una fila bajo el encabezado."""
    destino = tmp_path / "placas.xlsx"
    escribir_listado(filas, destino, directorio_imagenes)

    hoja = load_workbook(destino)["Placas"]

    assert hoja.cell(row=1, column=1).value == "Registro"
    assert hoja.max_row == 1 + len(filas)
    assert hoja.cell(row=2, column=2).value == "PCW2497"


def test_incrusta_la_imagen_de_respaldo(tmp_path, filas, directorio_imagenes):
    """La foto queda dentro del archivo, no como referencia externa."""
    destino = tmp_path / "placas.xlsx"
    escribir_listado(filas, destino, directorio_imagenes)

    hoja = load_workbook(destino)["Placas"]

    assert len(hoja._images) == 1


def test_fila_sin_imagen_no_rompe_el_listado(tmp_path, filas, directorio_imagenes):
    """Un vehículo sin placa se escribe igual, solo que sin foto."""
    destino = tmp_path / "placas.xlsx"
    escribir_listado(filas, destino, directorio_imagenes)

    hoja = load_workbook(destino)["Placas"]

    assert hoja.cell(row=3, column=1).value == "V002"


def test_imagen_inexistente_se_omite_sin_fallar(tmp_path, filas):
    """Una ruta de imagen rota no debe impedir la entrega del listado."""
    destino = tmp_path / "placas.xlsx"

    escribir_listado(filas, destino, tmp_path / "no_existe")

    assert destino.is_file()


def test_agrega_hoja_de_resumen(tmp_path, filas, directorio_imagenes):
    """El informe global se entrega en una hoja aparte."""
    destino = tmp_path / "placas.xlsx"
    resumen = {"Videos procesados": 65, "Vehiculos registrados": 412}

    escribir_listado(filas, destino, directorio_imagenes, resumen=resumen)

    libro = load_workbook(destino)
    assert "Resumen" in libro.sheetnames
    assert libro["Resumen"].cell(row=3, column=2).value == 65


def test_rechaza_extension_incorrecta(tmp_path, filas):
    """Solo se acepta escribir en formato xlsx."""
    with pytest.raises(ExcelError, match="xlsx"):
        escribir_listado(filas, tmp_path / "placas.csv")


def test_listado_vacio_produce_solo_encabezado(tmp_path):
    """Un lote sin vehículos genera el archivo con su encabezado."""
    destino = tmp_path / "vacio.xlsx"

    escribir_listado([], destino)

    assert load_workbook(destino)["Placas"].max_row == 1
