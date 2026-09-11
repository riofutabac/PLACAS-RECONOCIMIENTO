"""Generación del listado final en Excel con la foto de respaldo incrustada.

El CSV obliga a abrir cada imagen por separado para verificar una placa. Con
la foto dentro de la fila, revisar un registro dudoso es mirar la celda de al
lado, que es como el analista realmente trabaja.
"""

from dataclasses import dataclass
from pathlib import Path
from typing import Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ImagenExcel
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


class ExcelError(ValueError):
    """Excepción lanzada cuando el listado no puede generarse."""


# Alto de fila y ancho de columna pensados para que la foto se lea sin ampliar
ALTO_FILA_PUNTOS = 90
ALTO_IMAGEN_PIXELES = 115
ANCHO_COLUMNA_IMAGEN = 32

COLUMNAS = (
    ("registro", "Registro", 11),
    ("placa", "Placa", 13),
    ("hora_paso", "Hora de paso", 14),
    ("tiempo_video", "Tiempo en video", 15),
    ("video", "Archivo de video", 34),
    ("tipo_vehiculo", "Tipo", 12),
    ("sentido", "Sentido", 11),
    ("confianza", "Confianza", 11),
    ("consenso", "Consenso", 11),
    ("lecturas", "Lecturas", 10),
    ("estado", "Estado", 22),
)

COLOR_ENCABEZADO = "1F3864"
COLOR_VALIDADO = "E2EFDA"
COLOR_PENDIENTE = "FFF2CC"
COLOR_SIN_PLACA = "FCE4E4"

ESTADO_VALIDADO = "validado"
ESTADO_SIN_PLACA = "sin placa identificable"


def _color_de_fila(estado: str) -> str:
    """Color de fondo según el estado, para localizar los pendientes de un vistazo."""
    if estado == ESTADO_VALIDADO:
        return COLOR_VALIDADO
    if estado == ESTADO_SIN_PLACA:
        return COLOR_SIN_PLACA
    return COLOR_PENDIENTE


def _escribir_encabezado(hoja) -> None:
    """Escribe la fila de títulos y fija su formato."""
    relleno = PatternFill("solid", fgColor=COLOR_ENCABEZADO)
    for indice, (_, titulo, ancho) in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(row=1, column=indice, value=titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(horizontal="center", vertical="center")
        hoja.column_dimensions[get_column_letter(indice)].width = ancho

    columna_imagen = len(COLUMNAS) + 1
    celda = hoja.cell(row=1, column=columna_imagen, value="Imagen de respaldo")
    celda.font = Font(bold=True, color="FFFFFF")
    celda.fill = relleno
    celda.alignment = Alignment(horizontal="center", vertical="center")
    hoja.column_dimensions[get_column_letter(columna_imagen)].width = ANCHO_COLUMNA_IMAGEN
    hoja.freeze_panes = "A2"


def _insertar_imagen(hoja, ruta: Path, fila: int, columna: int) -> bool:
    """Incrusta la imagen en la celda, escalada al alto de la fila."""
    if not ruta.is_file():
        return False
    try:
        imagen = ImagenExcel(str(ruta))
    except Exception:
        return False

    proporcion = imagen.width / imagen.height if imagen.height else 1.0
    imagen.height = ALTO_IMAGEN_PIXELES
    imagen.width = int(ALTO_IMAGEN_PIXELES * proporcion)
    hoja.add_image(imagen, f"{get_column_letter(columna)}{fila}")
    return True


def escribir_listado(
    filas: Sequence[dict],
    ruta_salida,
    directorio_imagenes=None,
    resumen: Optional[dict] = None,
) -> Path:
    """Escribe el listado en Excel con la foto de cada registro incrustada.

    `directorio_imagenes` es la raíz contra la que se resuelven las rutas
    relativas de imagen de cada fila.
    """
    ruta = Path(ruta_salida)
    if ruta.suffix.lower() != ".xlsx":
        raise ExcelError(f"El listado debe escribirse en un archivo .xlsx, se recibió: '{ruta.name}'")

    raiz = Path(directorio_imagenes) if directorio_imagenes else ruta.parent
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Placas"
    _escribir_encabezado(hoja)

    columna_imagen = len(COLUMNAS) + 1
    for numero, datos in enumerate(filas, start=2):
        estado = str(datos.get("estado", ""))
        relleno = PatternFill("solid", fgColor=_color_de_fila(estado))

        for indice, (clave, _, _) in enumerate(COLUMNAS, start=1):
            celda = hoja.cell(row=numero, column=indice, value=datos.get(clave, ""))
            celda.fill = relleno
            celda.alignment = Alignment(horizontal="center", vertical="center")

        hoja.row_dimensions[numero].height = ALTO_FILA_PUNTOS
        ruta_imagen = datos.get("imagen", "")
        if ruta_imagen:
            _insertar_imagen(hoja, raiz / ruta_imagen, numero, columna_imagen)

    if resumen:
        _escribir_resumen(libro, resumen)

    libro.save(ruta)
    return ruta


def _escribir_resumen(libro, resumen: dict) -> None:
    """Agrega una hoja con el informe global del procesamiento."""
    hoja = libro.create_sheet("Resumen")
    hoja.column_dimensions["A"].width = 38
    hoja.column_dimensions["B"].width = 22

    titulo = hoja.cell(row=1, column=1, value="Informe del procesamiento")
    titulo.font = Font(bold=True, size=14)

    fila = 3
    for clave, valor in resumen.items():
        etiqueta = hoja.cell(row=fila, column=1, value=clave)
        etiqueta.font = Font(bold=True)
        hoja.cell(row=fila, column=2, value=valor)
        fila += 1
